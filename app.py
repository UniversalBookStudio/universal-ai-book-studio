import streamlit as st
import json, re, html, zipfile, os, tempfile, xml.etree.ElementTree as ET
from io import BytesIO
from datetime import datetime
from xml.sax.saxutils import escape

VERSION = "0.5.1"
TYPES = ["Auto Detect","Business Book","Novel / Fiction","Children's Book","Workbook / Activity Book","Textbook / Study Notes","Cookbook","Religious / Spiritual","History / Mythology","Self-help","Biography","Poetry","Magazine","Product Catalog","Technical Manual","Travel Book","Other"]
LANGS = ["Hindi","English","Hinglish","Bengali","Other","Auto Detect"]
STYLES = ["Director decides automatically","Premium Literary","Minimal","Modern","Luxury","Educational","Illustrated","Devotional","Editorial / Magazine","Custom"]

st.set_page_config(page_title="Universal AI Book Studio", page_icon="📚", layout="wide")

def detect_type(text, selected):
    if selected != "Auto Detect": return selected
    t = text.lower()
    rules = [
        (["workbook","worksheet","activity book","अभ्यास","वर्कबुक"], "Workbook / Activity Book"),
        (["children","kids","preschool","बच्चों","बाल"], "Children's Book"),
        (["recipe","cookbook","रेसिपी","व्यंजन"], "Cookbook"),
        (["novel","fiction","उपन्यास","कहानी"], "Novel / Fiction"),
        (["business","sales","marketing","दुकानदार","बिक्री","व्यापार"], "Business Book"),
        (["history","mythology","इतिहास","महाभारत","रामायण"], "History / Mythology"),
        (["spiritual","religious","धार्मिक","आध्यात्मिक","भक्ति"], "Religious / Spiritual"),
    ]
    for keys, val in rules:
        if any(k in t for k in keys): return val
    return "Other / General Guide"

def make_plan(idea, source, mode, kind, language, audience, goal, style, constraints):
    detected = detect_type(idea + "\n" + source[:3000], kind)
    design = {
        "Business Book":"Practical editorial layout; checklists, examples and tables",
        "Novel / Fiction":"Literary typography, immersive chapter openings, restrained ornaments",
        "Children's Book":"Large readable type, illustration-led pages, age-appropriate density",
        "Workbook / Activity Book":"Clear instructions, activity panels and generous response space",
        "Textbook / Study Notes":"Strong heading hierarchy, tables, callouts and diagrams where useful",
        "Cookbook":"Recipe cards with separated ingredients, method and serving notes",
        "Religious / Spiritual":"Respectful, calm, legible design without decorative clutter",
        "History / Mythology":"Editorial historical style; separate established facts from interpretation",
        "Product Catalog":"Consistent product cards, image zones and specification hierarchy"
    }.get(detected, "Clean, readable, genre-appropriate layout")
    if style != "Director decides automatically": design = style
    stages = (["Preserve original source", "Identify structure and front/back matter", "Run language and structure QC",
               "Lock approved content unless explicit permission is given", "Choose design system", "Typeset and inspect pages",
               "Validate exports"] if mode == "Manuscript to Book" else
              ["Clarify reader promise and book brief", "Build outline", "Identify research gaps",
               "Draft and revise chapters", "Check continuity, language and factual claims",
               "Select design system", "Typeset and inspect pages", "Validate exports"])
    return {
        "version": VERSION, "created_at": datetime.now().isoformat(timespec="seconds"),
        "mode": mode, "brief": idea.strip() or "Source manuscript project", "book_type": detected,
        "language": language, "target_audience": audience or "Infer cautiously and mark assumptions",
        "publishing_goal": goal, "design_director": {
            "style": design, "page_count": "Content-led; no forced fixed page count",
            "content_lock": "Never silently rewrite approved source content",
            "font_rule": "Use fonts verified for all required language glyphs",
            "visual_qc": "Inspect exported pages; automated checks alone are not print certification"
        }, "production_stages": stages,
        "quality_gates": ["source integrity", "spelling and grammar", "heading hierarchy",
          "repetition and continuity", "fact-check flags", "font/glyph rendering",
          "margins and page breaks", "image resolution and placement", "export validity"],
        "constraints": constraints,
        "limitations": ["PDF visual typesetting and image generation are not included in this version",
                        "KDP compliance is not certified automatically"]
    }

def clean_book_title(idea, fallback="Untitled Book"):
    """Extract a useful title from a multi-line brief without retaining labels like 'किताब का विषय:'."""
    ignored = {"किताब का विषय", "विषय", "book idea", "book topic", "title", "शीर्षक", "उद्देश्य", "पाठक", "भाषा", "शैली"}
    for raw in (idea or "").splitlines():
        line = raw.strip().strip("#*-• \t")
        if not line:
            continue
        # Remove common labels while preserving text after the colon.
        if ":" in line or "：" in line:
            left, right = re.split(r"[:：]", line, maxsplit=1)
            if left.strip().lower() in ignored:
                line = right.strip()
        line = line.strip('"“”‘’ ')
        if not line:
            continue
        if line.rstrip(":：").strip().lower() in ignored:
            continue
        # Skip instruction labels and pick the first meaningful content line.
        if line.lower().startswith(("उद्देश्य", "पाठक", "भाषा", "शैली", "महत्वपूर्ण नियम", "objective", "audience", "language", "style")):
            continue
        return line[:120]
    return fallback

def normalized_heading(text):
    return re.sub(r"\s+", " ", (text or "").strip()).casefold()

def local_outline(idea, source, mode, kind, language, audience, constraints):
    title = clean_book_title(idea, "Untitled Book")
    chosen = detect_type(idea + source[:2000], kind)
    if mode == "Manuscript to Book" and source.strip():
        headings = [x.strip().lstrip("# ").strip() for x in source.splitlines() if x.strip().startswith("#")]
        names = headings or ["Original manuscript — preserve source"]
        goal = "Structure extracted from source; original content remains separate and unchanged."
    else:
        names = ["Reader problem and desired outcome","Core concepts and foundations","Step-by-step method",
                 "Practical examples and templates","Common mistakes and solutions","Action plan and next steps"]
        goal = "Starter outline only; full chapter text requires an AI provider or manual writing."
    return {"title":title,"subtitle":"Working subtitle — refine after reviewing the brief",
      "book_type":chosen,"language":language,"target_readers":audience or "To be confirmed",
      "introduction_goal":goal,
      "chapters":[{"number":i+1,"title":name,"purpose":"Define reader outcome and supporting sections",
                   "sections":["Key idea","Explanation","Example or exercise","Chapter summary"]} for i,name in enumerate(names)],
      "research_tasks":["Verify factual claims and sources","Mark assumptions; do not invent citations"],
      "content_lock_rules":constraints or "Preserve approved text; rewrite only with explicit permission."}

def gemini_text(key, prompt, temperature=0.4):
    import requests
    r = requests.post("https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent",
      params={"key":key}, json={"contents":[{"parts":[{"text":prompt}]}],
      "generationConfig":{"temperature":temperature}}, timeout=120)
    r.raise_for_status()
    data = r.json()
    return data["candidates"][0]["content"]["parts"][0]["text"].strip()

def gemini_json(key, prompt):
    raw = gemini_text(key, prompt + "\nReturn valid JSON only, without markdown fences.")
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip(), flags=re.I)
    a,b = raw.find("{"),raw.rfind("}")
    if a < 0 or b < a: raise ValueError("AI output did not contain valid JSON.")
    return json.loads(raw[a:b+1])

def word_count(s): return len(re.findall(r"\b[\w'-]+\b", s, flags=re.UNICODE))

def build_md(title, outline, drafts):
    lines = ["# " + (title or outline.get("title","Untitled Book")), "", "_" + outline.get("subtitle","") + "_", ""]
    for h,body in drafts.items(): lines += ["## "+h, "", body.strip(), ""]
    return "\n".join(lines).strip()+"\n"

def build_html(title, outline, drafts):
    sections=[]
    for h,body in drafts.items():
        paras=[]
        for line in body.splitlines():
            line=line.strip()
            if not line: continue
            if line.startswith("### "): paras.append("<h3>"+escape(line[4:])+"</h3>")
            elif line.startswith("## "): paras.append("<h2>"+escape(line[3:])+"</h2>")
            elif line.startswith("# "): paras.append("<h2>"+escape(line[2:])+"</h2>")
            else: paras.append("<p>"+escape(line)+"</p>")
        sections.append("<section><h1>"+escape(h)+"</h1>"+"".join(paras)+"</section>")
    return '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+escape(title)+'</title><style>body{max-width:800px;margin:2rem auto;padding:0 1rem;font:18px/1.7 Georgia,serif;color:#222}h1,h2,h3{line-height:1.25}section{margin-bottom:3rem}@media print{body{max-width:none;margin:0}section{page-break-before:always}}</style></head><body><h1>'+escape(title)+'</h1><p><em>'+escape(outline.get("subtitle",""))+'</em></p>'+"".join(sections)+"</body></html>"

def build_docx(title, outline, drafts):
    from docx import Document
    doc=Document()
    doc.add_heading(title or outline.get("title","Untitled Book"),0)
    if outline.get("subtitle"): doc.add_paragraph(outline["subtitle"])
    doc.add_paragraph("Language: "+str(outline.get("language",""))+" | Type: "+str(outline.get("book_type","")))
    for heading,body in drafts.items():
        doc.add_heading(heading,1)
        for line in body.splitlines():
            line=line.strip()
            if not line: continue
            if line.startswith("### "): doc.add_heading(line[4:],3)
            elif line.startswith("## "): doc.add_heading(line[3:],2)
            elif line.startswith("# "): doc.add_heading(line[2:],1)
            else: doc.add_paragraph(line)
    out=BytesIO(); doc.save(out); return out.getvalue()

def build_epub(title, outline, drafts):
    title=escape(title or outline.get("title","Untitled Book"))
    chapters=list(drafts.items()) or [("Draft","<p>No chapter drafts have been added.</p>")]
    files=[]; manifest=[]; spine=[]; nav=[]
    for i,(heading,body) in enumerate(chapters,1):
        p=[]
        for line in body.splitlines():
            line=line.strip()
            if not line: continue
            if line.startswith("### "): p.append("<h3>"+escape(line[4:])+"</h3>")
            elif line.startswith("## "): p.append("<h2>"+escape(line[3:])+"</h2>")
            elif line.startswith("# "): p.append("<h2>"+escape(line[2:])+"</h2>")
            else: p.append("<p>"+escape(line)+"</p>")
        name=f"chapter{i}.xhtml"
        xhtml='<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>'+escape(heading)+'</title><meta charset="utf-8"/><link rel="stylesheet" href="style.css"/></head><body><h1>'+escape(heading)+'</h1>'+"".join(p)+"</body></html>"
        files.append((name,xhtml)); manifest.append(f'<item id="c{i}" href="{name}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="c{i}"/>'); nav.append(f'<li><a href="{name}">{escape(heading)}</a></li>')
    navx='<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents</title><meta charset="utf-8"/></head><body><nav epub:type="toc"><h1>Contents</h1><ol>'+"".join(nav)+'</ol></nav></body></html>'
    opf='<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">urn:uuid:universal-ai-book-studio</dc:identifier><dc:title>'+title+'</dc:title><dc:language>und</dc:language></metadata><manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/><item id="css" href="style.css" media-type="text/css"/>'+''.join(manifest)+'</manifest><spine>'+''.join(spine)+'</spine></package>'
    out=BytesIO()
    with zipfile.ZipFile(out,"w") as z:
        z.writestr("mimetype","application/epub+zip",compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",'<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        z.writestr("OEBPS/content.opf",opf); z.writestr("OEBPS/nav.xhtml",navx)
        z.writestr("OEBPS/style.css","body{font-family:serif;line-height:1.55;margin:5%}h1,h2,h3{line-height:1.2}p{margin:0 0 .8em}")
        for name,content in files: z.writestr("OEBPS/"+name,content)
    return out.getvalue()

def build_pdf(title, outline, drafts, font_bytes=None, font_name="BookUnicode"):
    """Basic PDF export. For Hindi/Bengali, upload a Unicode TTF font; shaping must be visually checked."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether
    from reportlab.lib import colors
    import os, tempfile

    tmp_path = None
    if font_bytes:
        with tempfile.NamedTemporaryFile(suffix=".ttf", delete=False) as f:
            f.write(font_bytes); tmp_path = f.name
        try:
            pdfmetrics.registerFont(TTFont(font_name, tmp_path))
        except Exception:
            os.unlink(tmp_path)
            raise ValueError("Font file पढ़ा नहीं जा सका। कृपया valid .ttf font upload करें।")
    else:
        font_name = "Helvetica"
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=22*mm, leftMargin=22*mm,
                            topMargin=22*mm, bottomMargin=20*mm, title=title,
                            author="Universal AI Book Studio")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BookTitle", parent=styles["Title"], fontName=font_name,
                              fontSize=22, leading=29, alignment=TA_CENTER, spaceAfter=18))
    styles.add(ParagraphStyle(name="BookSubtitle", parent=styles["Normal"], fontName=font_name,
                              fontSize=12, leading=18, alignment=TA_CENTER, spaceAfter=30))
    styles.add(ParagraphStyle(name="Chapter", parent=styles["Heading1"], fontName=font_name,
                              fontSize=17, leading=23, spaceBefore=8, spaceAfter=16, textColor=colors.HexColor("#243447")))
    styles.add(ParagraphStyle(name="Section2", parent=styles["Heading2"], fontName=font_name,
                              fontSize=13, leading=18, spaceBefore=12, spaceAfter=7))
    styles.add(ParagraphStyle(name="Section3", parent=styles["Heading3"], fontName=font_name,
                              fontSize=11, leading=15, spaceBefore=9, spaceAfter=5))
    styles.add(ParagraphStyle(name="BookBody", parent=styles["BodyText"], fontName=font_name,
                              fontSize=10.5, leading=16, spaceAfter=8, wordWrap="CJK"))
    story = [Spacer(1, 35*mm), Paragraph(html.escape(title or outline.get("title", "Untitled Book")), styles["BookTitle"])]
    subtitle = outline.get("subtitle", "")
    if subtitle: story.append(Paragraph(html.escape(subtitle), styles["BookSubtitle"]))
    story.append(Spacer(1, 12*mm))
    for heading, body in drafts.items():
        story.append(PageBreak())
        story.append(Paragraph(html.escape(heading), styles["Chapter"]))
        for line in body.splitlines():
            line=line.strip()
            if not line: continue
            if line.startswith("### "): story.append(Paragraph(html.escape(line[4:]), styles["Section3"]))
            elif line.startswith("## "): story.append(Paragraph(html.escape(line[3:]), styles["Section2"]))
            elif line.startswith("# "): story.append(Paragraph(html.escape(line[2:]), styles["Section2"]))
            else: story.append(Paragraph(html.escape(line).replace("\n", "<br/>"), styles["BookBody"]))
    def footer(canvas, doc_obj):
        canvas.saveState(); canvas.setFont(font_name, 8)
        canvas.drawCentredString(A4[0]/2, 10*mm, str(doc_obj.page))
        canvas.restoreState()
    try:
        doc.build(story, onFirstPage=footer, onLaterPages=footer)
    finally:
        if tmp_path and os.path.exists(tmp_path): os.unlink(tmp_path)
    return buf.getvalue()


def validate_epub(epub_bytes):
    """Structural EPUB package check; this is not a substitute for EPUBCheck or reader testing."""
    problems=[]
    try:
        with zipfile.ZipFile(BytesIO(epub_bytes)) as z:
            names=set(z.namelist())
            required={"mimetype","META-INF/container.xml","OEBPS/content.opf","OEBPS/nav.xhtml"}
            missing=required-names
            if missing: problems.append("Missing EPUB package files: " + ", ".join(sorted(missing)))
            if "mimetype" in names and z.read("mimetype") != b"application/epub+zip":
                problems.append("EPUB mimetype entry is incorrect.")
            for name in names:
                if name.endswith((".xml", ".xhtml", ".opf")):
                    try: ET.fromstring(z.read(name))
                    except Exception: problems.append("Invalid XML/XHTML: " + name)
    except Exception as e:
        problems.append("EPUB ZIP could not be read: " + str(e))
    return problems

def qc(title, drafts):
    errors=[]; warnings=[]
    if not title.strip(): errors.append("Book title is blank.")
    if not drafts: warnings.append("No chapter drafts yet.")
    seen={}
    for h,body in drafts.items():
        if not body.strip(): errors.append("Empty chapter: "+h)
        if word_count(body)<80: warnings.append(f"Short draft ({word_count(body)} words): {h}")
        if "[VERIFY]" in body: warnings.append("Fact-check marked claims in: "+h)
        if "TODO" in body or "[ADD " in body: warnings.append("Possible unfinished placeholder in: "+h)
        k=re.sub(r"\W+","",h.lower()); seen[k]=seen.get(k,0)+1
    for h,n in seen.items():
        if n>1: warnings.append("Duplicate chapter heading: "+h)
    return errors,warnings

st.title("📚 Universal AI Book Studio")
st.caption(f"Mobile-first • v{VERSION} • Plan → Outline → Drafts → QC → Export")
with st.sidebar:
    st.header("Project settings")
    mode=st.radio("Creation mode",["Idea to Book","Manuscript to Book","Idea + Partial Content"])
    kind=st.selectbox("Book type",TYPES)
    language=st.selectbox("Language",LANGS,index=0)
    audience=st.text_input("Target readers (optional)")
    goal=st.selectbox("Publishing goal",["Digital eBook","Print book","Amazon KDP","Web/interactive book","Multiple outputs"])
    style=st.selectbox("Design style",STYLES)
    api_key=st.text_input("Optional Gemini API key",type="password",help="API is unchanged in this version. Never paste the key into GitHub.")
    st.caption("API को छुए बिना local editing, QC और export काम कर सकते हैं।")

st.subheader("PDF font (optional)")
pdf_font = st.file_uploader("Hindi/Bengali PDF के लिए Unicode .TTF font upload करें", type=["ttf"], help="उदाहरण: Noto Sans Devanagari का .ttf फ़ॉन्ट। बिना Unicode font के Hindi PDF सही नहीं दिख सकती।")
if pdf_font:
    st.success("PDF font loaded for this session. इसे GitHub पर upload नहीं किया जाएगा।")

a,b=st.columns([1.1,0.9])
with a:
    idea=st.text_area("Book idea / brief",height=120,placeholder="किताब का विषय और पाठक का लक्ष्य लिखें।")
    upload=st.file_uploader("Optional manuscript/notes (TXT, MD)",type=["txt","md"])
    source=upload.getvalue().decode("utf-8",errors="replace") if upload else ""
    if upload: st.success(f"{upload.name} loaded • {len(source):,} characters")
    constraints=st.text_area("Special instructions / content locks",height=80,placeholder="स्वीकृत टेक्स्ट बिना अनुमति न बदलें।")
with b:
    st.subheader("Master Design Director")
    st.markdown("- Genre-aware production planning\n- No forced page count\n- Editable outline and chapters\n- Preliminary content QC\n- EPUB / DOCX / HTML / MD / TXT exports")
    st.info("PDF export उपलब्ध है, लेकिन अभी basic है। हिंदी/बंगाली के लिए सही Unicode TTF फ़ॉन्ट दें। वास्तविक पेज-दर-पेज विज़ुअल QC, image generation और KDP certification अभी उपलब्ध नहीं हैं।")

if st.button("Create / refresh production plan",type="primary",use_container_width=True):
    if not idea.strip() and not source.strip(): st.warning("Idea लिखें या manuscript upload करें।")
    else:
        st.session_state.plan=make_plan(idea,source,mode,kind,language,audience,goal,style,constraints)
        st.session_state.inputs={"idea":idea,"source":source,"mode":mode,"kind":kind,"language":language,"audience":audience,"goal":goal,"style":style,"constraints":constraints}
        st.success("Production plan तैयार है।")
if "plan" in st.session_state:
    with st.expander("Production Plan JSON"):
        st.json(st.session_state.plan)
    st.download_button("Download plan JSON",json.dumps(st.session_state.plan,ensure_ascii=False,indent=2),"production_plan.json","application/json")

st.divider(); st.header("✍️ Manuscript Workshop")
x,y=st.columns(2)
with x:
    if st.button("Build starter outline (local)",use_container_width=True):
        if not idea.strip() and not source.strip(): st.warning("पहले idea लिखें या source upload करें।")
        else:
            st.session_state.outline=local_outline(idea,source,mode,kind,language,audience,constraints)
            st.success("Starter outline तैयार है।")
with y:
    if st.button("Build outline with Gemini AI",disabled=not bool(api_key.strip()),use_container_width=True):
        try:
            prompt=f"""Create JSON keys title,subtitle,book_type,language,target_readers,introduction_goal,chapters (objects number,title,purpose,sections),research_tasks,content_lock_rules. No forced page count or filler. Preserve source content for Manuscript to Book. IDEA:{idea} MODE:{mode} TYPE:{kind} LANGUAGE:{language} AUDIENCE:{audience} CONSTRAINTS:{constraints} SOURCE:{source[:12000]}"""
            st.session_state.outline=gemini_json(api_key.strip(),prompt); st.success("AI outline तैयार है।")
        except Exception as e: st.error(f"AI outline नहीं बन सका: {e}")

if "outline" in st.session_state:
    o=st.session_state.outline
    st.subheader("Outline — review and edit")
    o["title"]=st.text_input("Book title",o.get("title","Untitled Book"),key="outline_title")
    o["subtitle"]=st.text_input("Subtitle",o.get("subtitle",""),key="outline_subtitle")
    chs=o.setdefault("chapters",[])
    for i,ch in enumerate(chs):
        if not isinstance(ch,dict): continue
        with st.expander(f"Chapter {i+1}: {ch.get('title','Untitled')}"):
            ch["title"]=st.text_input("Chapter title",ch.get("title",""),key=f"ch_title_{i}")
            ch["purpose"]=st.text_area("Purpose",ch.get("purpose",""),key=f"ch_purpose_{i}",height=65)
            ch["sections"]= [s.strip() for s in st.text_area("Sections (one per line)","\n".join(ch.get("sections",[])),key=f"ch_sections_{i}",height=90).splitlines() if s.strip()]
    q1,q2=st.columns(2)
    with q1:
        if st.button("Add chapter"):
            chs.append({"number":len(chs)+1,"title":"New chapter","purpose":"","sections":[]}); st.rerun()
    with q2:
        if chs:
            remove=st.selectbox("Chapter to remove",[f"{i+1}. {c.get('title','Untitled')}" for i,c in enumerate(chs)],key="remove_chapter")
            if st.button("Remove selected chapter"):
                del chs[int(remove.split(".")[0])-1]; st.rerun()
    st.download_button("Download outline JSON",json.dumps(o,ensure_ascii=False,indent=2),"book_outline.json","application/json")
    if chs:
        selected=st.selectbox("Select chapter to draft",range(len(chs)),format_func=lambda i:f"{i+1}. {chs[i].get('title','Untitled')}")
        ch=chs[selected]
        if st.button("Draft selected chapter with Gemini",disabled=not bool(api_key.strip()),type="primary",use_container_width=True):
            lock="Preserve source manuscript exactly; do not silently rewrite it." if mode=="Manuscript to Book" else "Do not invent citations, statistics or sources. Mark uncertain claims [VERIFY]."
            prompt=f"Write one full chapter in {language}. Clear headings and examples. Return only chapter text. OUTLINE:{json.dumps(o,ensure_ascii=False)} CHAPTER:{json.dumps(ch,ensure_ascii=False)} RULES:{constraints} LOCK:{lock} SOURCE:{source[:10000]}"
            try:
                with st.spinner("Drafting chapter..."):
                    draft=gemini_text(api_key.strip(),prompt,0.45)
                st.session_state.setdefault("drafts",{})[ch.get("title",f"Chapter {selected+1}")]=draft
                st.success("Draft तैयार है; publication से पहले review करें।")
            except Exception as e: st.error(f"Draft नहीं बन सका: {e}")

st.divider(); st.subheader("Add or edit chapter text manually")
st.caption("Chapter heading = अध्याय का शीर्षक। Chapter text = उस अध्याय की पूरी सामग्री। अगर Outline में अध्याय पहले से है, तो उसी अध्याय को चुनें ताकि duplicate chapter न बने।")
outline_for_manual = st.session_state.get("outline", {})
outline_chapters_for_manual = outline_for_manual.get("chapters", []) if isinstance(outline_for_manual, dict) else []
manual_choices = ["नया अध्याय बनाएँ"] + [
    f"{i+1}. {c.get('title','Untitled')}" for i,c in enumerate(outline_chapters_for_manual) if isinstance(c, dict)
]
manual_target = st.selectbox("यह टेक्स्ट किस अध्याय का है?", manual_choices, key="manual_target")
default_manual_title = ""
if manual_target != "नया अध्याय बनाएँ":
    try:
        manual_index = int(manual_target.split(".", 1)[0]) - 1
        default_manual_title = outline_chapters_for_manual[manual_index].get("title", "")
    except Exception:
        default_manual_title = ""
manual_title = st.text_input("Chapter heading — अध्याय का शीर्षक", value=default_manual_title, key="manual_title_v051")
manual_text = st.text_area("Chapter text — अध्याय की पूरी सामग्री", height=170, key="manual_text_v051",
                           placeholder="यहाँ अध्याय की सामग्री लिखें या paste करें।")
if st.button("Save / update chapter text", use_container_width=True):
    if not manual_title.strip() or not manual_text.strip():
        st.warning("अध्याय का शीर्षक और पूरी सामग्री—दोनों भरें।")
    else:
        st.session_state.setdefault("drafts", {})[manual_title.strip()] = manual_text
        st.success(f"‘{manual_title.strip()}’ का टेक्स्ट सेव हो गया।")

if "drafts" in st.session_state and st.session_state.drafts:
    st.divider(); st.header("🔎 Quality Check and Export")
    drafts=st.session_state.drafts
    for heading in list(drafts.keys()):
        with st.expander(heading,expanded=False):
            drafts[heading]=st.text_area("Edit chapter",drafts[heading],height=240,key="draft_edit_"+heading)
            if st.button("Delete this chapter",key="draft_del_"+heading):
                del drafts[heading]; st.rerun()
    outline=st.session_state.get("outline",{})
    default_export_title = outline.get("title", "Untitled Book")
    if default_export_title.strip().rstrip(":：").lower() in {"किताब का विषय", "विषय", "book idea", "book topic", "title", "शीर्षक"}:
        default_export_title = clean_book_title(st.session_state.get("inputs", {}).get("idea", ""), "Untitled Book")
    title=st.text_input("Export title — अंतिम किताब का शीर्षक",default_export_title,key="export_title")
    errors,warnings=qc(title,drafts)
    total=sum(word_count(v) for v in drafts.values())
    m1,m2=st.columns(2); m1.metric("Total words",f"{total:,}"); m2.metric("Draft chapters",len(drafts))
    for e in errors: st.error(e)
    if not errors: st.success("No blocking issue found by the local checks.")
    for w in warnings: st.warning(w)
    st.caption("यह केवल शुरुआती स्वचालित जाँच है। यह हिंदी व्याकरण, तथ्य-सत्यता या संदर्भों को प्रमाणित नहीं करती; प्रकाशन से पहले मानवीय समीक्षा आवश्यक है।")
    md=build_md(title,outline,drafts); html_doc=build_html(title,outline,drafts)
    package={"version":VERSION,"saved_at":datetime.now().isoformat(timespec="seconds"),"inputs":st.session_state.get("inputs",{}),"outline":outline,"drafts":drafts,"qc":{"errors":errors,"warnings":warnings}}
    c1,c2,c3=st.columns(3)
    with c1:
        try:
            if language in ["Hindi", "Bengali"] and not pdf_font:
                st.caption("हिंदी/बंगाली PDF के लिए पहले Unicode .TTF font upload करें।")
            pdf_bytes = build_pdf(title, outline, drafts, pdf_font.getvalue() if pdf_font else None)
            st.download_button("Download PDF", pdf_bytes, "book.pdf", "application/pdf", use_container_width=True)
        except Exception as e:
            st.warning(f"PDF export नहीं बन सका: {e}")
        st.download_button("Download Markdown",md,"book.md","text/markdown",use_container_width=True)
        st.download_button("Download TXT",md,"book.txt","text/plain",use_container_width=True)
    with c2:
        st.download_button("Download HTML",html_doc,"book.html","text/html",use_container_width=True)
        try: st.download_button("Download DOCX",build_docx(title,outline,drafts),"book.docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True)
        except Exception as e: st.warning(f"DOCX export error: {e}")
    with c3:
        try:
            epub_bytes = build_epub(title,outline,drafts)
            epub_problems = validate_epub(epub_bytes)
            st.download_button("Download EPUB",epub_bytes,"book.epub","application/epub+zip",use_container_width=True)
            if epub_problems:
                for problem in epub_problems: st.warning(problem)
            else:
                st.success("EPUB package structure check passed. Reader compatibility and visual formatting still need separate testing.")
        except Exception as e: st.warning(f"EPUB export error: {e}")
        st.download_button("Download project backup",json.dumps(package,ensure_ascii=False,indent=2),"book_project_backup.json","application/json",use_container_width=True)

st.divider()
st.subheader("Production status")
st.markdown("""
- **Available locally:** production plan, starter outline, manual writing/editing, outline edits, local QC, JSON backup, MD/TXT/HTML/DOCX/EPUB exports.
- **Still pending:** advanced page templates, image generation, page-by-page visual QC, automated KDP preflight and deeper content-lock enforcement. PDF export is basic and needs visual review; Hindi/Bengali require a Unicode font upload.
- **API:** unchanged by this update, as requested.
""")
st.caption(f"Universal AI Book Studio v{VERSION} • Keep API keys out of public GitHub.")
